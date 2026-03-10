"""Career report service."""

from __future__ import annotations

import json
import re
from io import BytesIO
from typing import Any

from ..db import ensure_business_tables, get_db_conn, to_iso
from .job_service import get_job_profile
from .matching_service import build_matching_analysis
from .student_service import get_student_profile


def generate_report(
    student_id: int,
    job_id: int,
    custom_instructions: str | None = None,
    user_id: str = "default",
) -> dict[str, Any]:
    student = get_student_profile(student_id)
    job = get_job_profile(job_id)
    analysis = build_matching_analysis(student["profile"], job)

    title = f"{student.get('student_name') or '学生'} - {job.get('title') or '目标岗位'} 职业规划报告"
    sections = [
        {
            "title": "职业探索与岗位匹配",
            "content": f"目标岗位为{job.get('title')}，当前人岗匹配总分{analysis['overall_score']}分。",
        },
        {
            "title": "能力差距分析",
            "content": "；".join([f"{x['name']}差值{x['gap']}" for x in analysis["dimensions"]]),
        },
        {
            "title": "行动计划",
            "content": "先补齐关键缺失技能，再通过项目与实习强化岗位竞争力。",
        },
    ]
    if custom_instructions:
        sections.append({"title": "个性化补充", "content": custom_instructions})

    content = {"title": title, "sections": sections}

    conn = get_db_conn()
    try:
        ensure_business_tables(conn)
        with conn.cursor() as cur:
            cur.execute(
                """
                INSERT INTO career_report (student_id, job_id, status, content)
                VALUES (%s, %s, %s, %s)
                RETURNING id
                """,
                (student_id, job_id, "ready", json.dumps(content, ensure_ascii=False)),
            )
            report_id = cur.fetchone()[0]

            cur.execute(
                """
                INSERT INTO matching_history (user_id, student_id, job_id, overall_score, analysis, report_id)
                VALUES (%s, %s, %s, %s, %s, %s)
                """,
                (
                    user_id or "default",
                    student_id,
                    job_id,
                    analysis["overall_score"],
                    json.dumps(analysis, ensure_ascii=False),
                    report_id,
                ),
            )
        conn.commit()
    finally:
        conn.close()

    return {"report_id": report_id, "status": "ready"}


def get_report(report_id: int) -> dict[str, Any]:
    """Fetch report by id."""
    conn = get_db_conn()
    try:
        ensure_business_tables(conn)
        with conn.cursor() as cur:
            cur.execute(
                "SELECT id, status, content, created_at FROM career_report WHERE id = %s",
                (report_id,),
            )
            row = cur.fetchone()
            if not row:
                raise ValueError("报告不存在")
            return {
                "report_id": row[0],
                "status": row[1],
                "content": row[2],
                "created_at": to_iso(row[3]),
            }
    finally:
        conn.close()


def polish_report(report_id: int, paragraph_index: int, text: str) -> dict[str, Any]:
    """Simple in-place paragraph polishing for report sections."""
    polished = re.sub(r"\s+", " ", (text or "").strip())
    if not polished:
        polished = "内容为空，建议补充具体行动目标。"

    conn = get_db_conn()
    try:
        ensure_business_tables(conn)
        with conn.cursor() as cur:
            cur.execute("SELECT content FROM career_report WHERE id = %s", (report_id,))
            row = cur.fetchone()
            if not row:
                raise ValueError("报告不存在")

            content = row[0] or {}
            sections = content.get("sections") or []
            if 0 <= paragraph_index < len(sections):
                sections[paragraph_index]["content"] = polished
                content["sections"] = sections
                cur.execute(
                    "UPDATE career_report SET content = %s, updated_at = CURRENT_TIMESTAMP WHERE id = %s",
                    (json.dumps(content, ensure_ascii=False), report_id),
                )
        conn.commit()
    finally:
        conn.close()

    return {"polished_text": polished}


def export_report(report_id: int, fmt: str) -> tuple[str, bytes, str]:
    """Export report as real pdf/docx/txt binary payload."""
    report = get_report(report_id)
    content = report.get("content") or {}

    lines = [content.get("title", f"report-{report_id}")]
    for sec in content.get("sections", []):
        lines.append(f"\n## {sec.get('title', '')}\n{sec.get('content', '')}")
    text = "\n".join(lines)

    fmt = (fmt or "pdf").lower()
    if fmt not in {"pdf", "docx", "txt"}:
        fmt = "txt"

    mime_map = {
        "pdf": "application/pdf",
        "docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "txt": "text/plain; charset=utf-8",
    }

    if fmt == "docx":
        try:
            from docx import Document
        except ImportError as e:
            raise ValueError("缺少 python-docx 依赖，无法导出 DOCX") from e

        doc = Document()
        doc.add_heading(content.get("title", f"report-{report_id}"), level=1)
        for sec in content.get("sections", []):
            doc.add_heading(sec.get("title", ""), level=2)
            doc.add_paragraph(sec.get("content", ""))

        output = BytesIO()
        doc.save(output)
        data = output.getvalue()
    elif fmt == "pdf":
        try:
            from reportlab.lib.pagesizes import A4
            from reportlab.pdfbase import pdfmetrics
            from reportlab.pdfbase.cidfonts import UnicodeCIDFont
            from reportlab.pdfgen import canvas
        except ImportError as e:
            raise ValueError("缺少 reportlab 依赖，无法导出 PDF") from e

        output = BytesIO()
        c = canvas.Canvas(output, pagesize=A4)
        _, height = A4
        margin = 40
        y = height - margin

        # Use a CJK CID font to avoid Chinese glyph squares.
        font_name = "STSong-Light"
        pdfmetrics.registerFont(UnicodeCIDFont(font_name))

        c.setFont(font_name, 16)
        for line in str(content.get("title", f"report-{report_id}")).splitlines() or [""]:
            c.drawString(margin, y, line)
            y -= 24
            if y < margin:
                c.showPage()
                c.setFont(font_name, 16)
                y = height - margin

        c.setFont(font_name, 12)
        for sec in content.get("sections", []):
            section_title = f"## {sec.get('title', '')}"
            for line in section_title.splitlines() or [""]:
                c.drawString(margin, y, line)
                y -= 18
                if y < margin:
                    c.showPage()
                    c.setFont(font_name, 12)
                    y = height - margin

            for line in str(sec.get("content", "")).splitlines() or [""]:
                c.drawString(margin, y, line)
                y -= 16
                if y < margin:
                    c.showPage()
                    c.setFont(font_name, 12)
                    y = height - margin

            y -= 8
            if y < margin:
                c.showPage()
                c.setFont(font_name, 12)
                y = height - margin

        c.save()
        data = output.getvalue()
    else:
        data = text.encode("utf-8")

    filename = f"report_{report_id}.{fmt}"
    return filename, data, mime_map[fmt]

